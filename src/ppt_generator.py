"""PPTX creation engine — safe layouts via built-in templates + text formatting."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from io import BytesIO

from src.templates import (
    THEMES, FONT_FAMILIES, SLIDE_WIDTH, SLIDE_HEIGHT,
)
from src.models import PPTRequest
from src.image_service import get_picsum_bg, get_picsum_square


# ============================================================================
# MAIN ENTRY
# ============================================================================
def create_pptx(slide_data: dict, request: PPTRequest) -> BytesIO:
    theme = THEMES[request.style.value]

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slides = slide_data.get("slides", [])
    for i, info in enumerate(slides):
        st = info.get("type", "content")
        depth = info.get("depth", "medium")

        if st == "title":
            _title(prs, info, theme)
        elif st == "section":
            _section(prs, info, theme)
        elif st == "chart":
            _chart(prs, info, theme)
        elif st == "two_column":
            _two_col(prs, info, theme)
        elif st == "quote":
            _section(prs, info, theme)
        elif st == "image_slide":
            _image_slide(prs, info, theme)
        elif st == "timeline":
            _timeline(prs, info, theme)
        elif st == "comparison":
            _comparison(prs, info, theme)
        elif st == "big_idea":
            _big_idea(prs, info, theme)
        else:
            _content(prs, info, theme, depth)

        # Decorations — applied to every slide (safe shape-based only)
        slide = prs.slides[-1]

        # Top accent bar (except title slide which has its own design)
        if i > 0 and st not in ("image_slide", "big_idea"):
            _top_bar(slide, theme)
            _slide_num(slide, i, theme)

        # Bottom line on content-type slides
        if st in ("content", "agenda", "summary", "chart", "two_column", "datapoint"):
            _bottom_bar(slide, theme)

        # Section slides: left accent stripe
        if st == "section":
            _left_stripe(slide, theme)

        # Slide progress indicator (skip title/section/big_idea)
        if i > 0 and st not in ("section", "image_slide", "big_idea"):
            _slide_progress(slide, i, len(slides), theme)

        _notes(slide, info)

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================================
# FORMAT HELPERS  (safe — only manipulate paragraph runs)
# ============================================================================

def _hex(h: str) -> RGBColor:
    """Convert hex string '1F4E79' to RGBColor."""
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _run(paragraph, text, size=None, color=None, bold=False, italic=False):
    """Add a run with formatted text. Safe — no XML manipulation."""
    run = paragraph.add_run()
    run.text = text
    if size:
        run.font.size = size
    if color:
        run.font.color.rgb = _hex(color)
    run.font.bold = bold
    run.font.italic = italic
    return run


def _thin_line(slide, left, top, width, color_hex):
    """Add a thin decorative line. Safe rectangle shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(0.015),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(color_hex)
    shape.line.fill.background()


# ============================================================================
# SLIDE LAYOUTS
# ============================================================================

def _title(prs, info, theme):
    """Title slide with accent line."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    title = info.get("title", "Presentation")
    subtitle = info.get("subtitle", "")

    # Style title shape
    if slide.shapes.title:
        tf = slide.shapes.title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        _run(p, title, size=Pt(36), color=theme.title_color, bold=True)
        p.alignment = PP_ALIGN.CENTER

    # Accent line under title
    _thin_line(slide, 4.5, 2.7, 4.333, theme.accent_bar_color)

    # Bottom decorative band
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(6.8), Inches(13.333), Inches(0.7),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = _hex(theme.primary_color)
    band.line.fill.background()

    # Subtitle
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            tf2 = ph.text_frame
            tf2.clear()
            p2 = tf2.paragraphs[0]
            _run(p2, subtitle, size=Pt(18), color="888888")
            p2.alignment = PP_ALIGN.CENTER
            break


def _section(prs, info, theme):
    """Section divider with accent lines above and below."""
    slide = prs.slides.add_slide(prs.slide_layouts[2])

    title = info.get("title", "")
    text = info.get("quote", "")

    # Accent line above
    _thin_line(slide, 5.0, 2.8, 3.333, theme.accent_bar_color)

    # Title
    display = title if title else text
    if slide.shapes.title:
        tf = slide.shapes.title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        _run(p, display, size=Pt(32), color="333333", bold=True)
        p.alignment = PP_ALIGN.CENTER

    # Accent line below (only for section, not quote)
    if title and not text:
        _thin_line(slide, 5.0, 4.8, 3.333, theme.accent_bar_color)

    # Attribution (for quote type)
    attr = info.get("attribution", "")
    if attr:
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                tf2 = ph.text_frame
                tf2.clear()
                p2 = tf2.paragraphs[0]
                _run(p2, f"— {attr}", size=Pt(14), color="999999", italic=True)
                break


def _content(prs, info, theme, depth="medium"):
    """Content slide — title + styled body with bullet prefixes."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    # === Title ===
    title = info.get("title", "")
    if slide.shapes.title:
        tf = slide.shapes.title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        _run(p, title, size=Pt(26), color=theme.title_color, bold=True)
        # Thin rule under title
        _thin_line(slide, 0.8, 1.05, 3.0, theme.accent_bar_color)

    # === Body ===
    body_ph = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            body_ph = ph
            break

    if not body_ph:
        return

    tf = body_ph.text_frame
    tf.clear()

    lines = []  # (text, size, color, bold, italic)

    # Subtitle
    sub = info.get("subtitle", "")
    if sub:
        lines.append((sub, Pt(14), theme.secondary_color, False, True))
        lines.append(("", Pt(6), "FFFFFF", False, False))  # spacer

    # Narrative (light slides)
    narrative = info.get("narrative", "")
    if narrative:
        lines.append((narrative, Pt(20) if depth == "light" else Pt(16), theme.text_color, False, False))
        lines.append(("", Pt(8), "FFFFFF", False, False))

    # Bullets
    bullets = info.get("bullets", [])
    for item in bullets:
        if isinstance(item, dict):
            main_text = item.get("text", "")
            subs = item.get("sub_bullets", [])
        else:
            main_text = str(item)
            subs = []

        if main_text:
            lines.append((f"  {main_text}", Pt(16), theme.text_color, False, False))
        for s_text in subs:
            lines.append((f"     ─ {s_text}", Pt(13), theme.secondary_color, False, False))

    # Highlight
    highlight = info.get("highlight", "")
    if highlight:
        lines.append(("", Pt(8), "FFFFFF", False, False))  # spacer
        lines.append((f"☆  {highlight}", Pt(15), theme.accent_bar_color, True, False))

    # Big number (datapoint)
    big = info.get("big_number", "")
    if big:
        lines.insert(0, (big, Pt(56), theme.accent_bar_color, True, False))
        desc = info.get("description", "")
        if desc:
            lines.insert(1, (desc, Pt(16), theme.text_color, False, False))

    # Write to placeholder
    for j, (text, size, color, bold, italic) in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        if text.strip():  # non-empty
            _run(p, text, size=size, color=color, bold=bold, italic=italic)
        else:  # spacer
            _run(p, "", size=Pt(4), color="FFFFFF")
        p.space_after = Pt(2)
        p.space_before = Pt(2)


def _two_col(prs, info, theme):
    """Two-column slide using built-in Two Content layout."""
    slide = prs.slides.add_slide(prs.slide_layouts[3])

    # Title
    title = info.get("title", "")
    if slide.shapes.title:
        tf = slide.shapes.title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        _run(p, title, size=Pt(26), color=theme.title_color, bold=True)
        _thin_line(slide, 0.8, 1.05, 3.0, theme.accent_bar_color)

    # Left column → idx 1
    left = info.get("left_bullets", [])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            tf = ph.text_frame
            tf.clear()
            for j, item in enumerate(left):
                text = item.get("text", "") if isinstance(item, dict) else str(item)
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                _run(p, f"  {text}", size=Pt(15), color=theme.text_color)
                p.space_after = Pt(6)
            break

    # Right column → idx 2
    right = info.get("right_bullets", [])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 2:
            tf = ph.text_frame
            tf.clear()
            for j, item in enumerate(right):
                text = item.get("text", "") if isinstance(item, dict) else str(item)
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                _run(p, f"  {text}", size=Pt(15), color=theme.text_color)
                p.space_after = Pt(6)
            break


def _chart(prs, info, theme):
    """Chart slide — blank layout + chart + styled insight panel."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title = info.get("title", "")
    if title:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11.5), Inches(0.6))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        _run(p, title, size=Pt(26), color=theme.title_color, bold=True)
        _thin_line(slide, 0.8, 0.95, 3.0, theme.accent_bar_color)

    # Subtitle
    sub = info.get("subtitle", "")
    chart_y = Inches(1.3) if sub else Inches(1.1)
    if sub:
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.5), Inches(0.3))
        tb2.text_frame.word_wrap = True
        p2 = tb2.text_frame.paragraphs[0]
        _run(p2, sub, size=Pt(14), color=theme.secondary_color, italic=True)

    # Build chart
    ct_str = info.get("chart_type", "column")
    chart_types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
        "area": XL_CHART_TYPE.AREA,
    }
    xl_type = chart_types.get(ct_str, XL_CHART_TYPE.COLUMN_CLUSTERED)

    cats = info.get("categories", [])
    series_list = info.get("series", [])
    cd = CategoryChartData()
    cd.categories = cats
    for s in series_list:
        cd.add_series(s.get("name", "Series"), s.get("values", []))

    chart_frame = slide.shapes.add_chart(
        xl_type, Inches(0.6), chart_y, Inches(7.8), Inches(5.2), cd,
    )
    chart = chart_frame.chart

    # Chart styling
    chart.chart_style = 3
    if ct_str in ("bar", "column", "stacked_bar"):
        try:
            chart.plots[0].gap_width = 80  # thicker bars
        except Exception:
            pass

    # Data labels
    if ct_str in ("bar", "column", "stacked_bar", "pie"):
        try:
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(8)
        except Exception:
            pass

    # Series colors
    accent_colors = [theme.accent_bar_color, theme.primary_color,
                     theme.secondary_color, theme.title_color]
    for idx, series in enumerate(chart.series):
        color = accent_colors[idx % len(accent_colors)]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _hex(color)

    # Line styling
    if ct_str in ("line", "area"):
        for idx, series in enumerate(chart.series):
            color = accent_colors[idx % len(accent_colors)]
            series.format.line.color.rgb = _hex(color)
            series.format.line.width = Pt(2.5)

    # Legend at bottom
    if chart.has_legend:
        chart.legend.font.size = Pt(9)
        try:
            from pptx.enum.chart import XL_LEGEND_POSITION
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        except ImportError:
            pass

    # Axis labels
    if ct_str != "pie":
        if chart.category_axis:
            chart.category_axis.tick_labels.font.size = Pt(8)
        if chart.value_axis:
            chart.value_axis.tick_labels.font.size = Pt(8)

    # Insight box
    insight = info.get("insight", "")
    if insight:
        # Subtle background card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(8.7), Inches(2.3), Inches(4.0), Inches(3.5),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = _hex(theme.primary_color)
        card.line.fill.background()

        # Label
        lb = slide.shapes.add_textbox(Inches(8.95), Inches(2.5), Inches(3.5), Inches(0.3))
        plb = lb.text_frame.paragraphs[0]
        _run(plb, "KEY INSIGHT", size=Pt(8), color="B0C4DE", bold=True)

        # Insight text
        ib = slide.shapes.add_textbox(Inches(8.95), Inches(2.85), Inches(3.5), Inches(2.7))
        ib.text_frame.word_wrap = True
        pib = ib.text_frame.paragraphs[0]
        _run(pib, insight, size=Pt(13), color="FFFFFF")


# ============================================================================
# NEW SLIDE TYPES — visual upgrades
# ============================================================================

def _image_slide(prs, info, theme):
    """Full-bleed image background with bold text overlay. Falls back to color bg."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Try to download a background image
    desc = info.get("image_description", info.get("title", "abstract"))
    img_buf = get_picsum_bg(desc, blur=0)

    if img_buf:
        # Add image as full-slide background
        slide.shapes.add_picture(img_buf, Inches(0), Inches(0),
                                 Inches(13.333), Inches(7.5))
        # Dark overlay for text readability
        overlay = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5),
        )
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = RGBColor(0, 0, 0)
        overlay.line.fill.background()
        # Set transparency via alpha
        try:
            from pptx.oxml.ns import qn
            sf = overlay.fill._fill
            srgb = sf.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            if srgb is not None:
                a = srgb.makeelement(qn('a:alpha'), {'val': '40000'})
                srgb.append(a)
        except Exception:
            pass  # transparency is optional
        text_color = "FFFFFF"
    else:
        # Fallback: solid color background
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5),
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = _hex(theme.primary_color)
        rect.line.fill.background()
        text_color = "FFFFFF"

    # Big statement text
    content = info.get("narrative") or info.get("title", "")
    tb = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.333), Inches(2.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    _run(p, content, size=Pt(32), color=text_color, bold=True)
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    sub = info.get("subtitle", "")
    if sub:
        tb2 = slide.shapes.add_textbox(Inches(2.0), Inches(5.0), Inches(9.333), Inches(0.8))
        tb2.text_frame.word_wrap = True
        p2 = tb2.text_frame.paragraphs[0]
        _run(p2, sub, size=Pt(16), color="B0C4DE" if img_buf else "D0D0D0")
        p2.alignment = PP_ALIGN.CENTER


def _timeline(prs, info, theme):
    """Horizontal timeline with milestone nodes."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    title = info.get("title", "Timeline")
    if slide.shapes.title:
        tf = slide.shapes.title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        _run(p, title, size=Pt(26), color=theme.title_color, bold=True)
        _thin_line(slide, 0.8, 1.05, 3.0, theme.accent_bar_color)

    milestones = info.get("milestones", [])
    if not milestones:
        return

    n = len(milestones)
    line_y = Inches(3.6)
    start_x = Inches(0.8)
    end_x = Inches(12.5)
    total_w = 11.7  # inches of usable space

    # Horizontal line
    _thin_line(slide, 0.8, 3.6, total_w, theme.accent_bar_color)

    for idx, m in enumerate(milestones):
        x = start_x + Emu(int(Emu(Inches(total_w)) * idx / max(n - 1, 1)))

        # Node circle
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            x - Inches(0.12), line_y - Inches(0.12), Inches(0.24), Inches(0.24),
        )
        node.fill.solid()
        node.fill.fore_color.rgb = _hex(theme.accent_bar_color)
        node.line.fill.background()

        # Year label
        year = m.get("year", "") if isinstance(m, dict) else str(m)
        event = m.get("title", "") if isinstance(m, dict) else ""
        brief = m.get("brief", "") if isinstance(m, dict) else ""

        if year:
            ty = slide.shapes.add_textbox(x - Inches(0.6), line_y - Inches(0.6), Inches(1.2), Inches(0.35))
            tp = ty.text_frame.paragraphs[0]
            _run(tp, str(year), size=Pt(12), color=theme.accent_bar_color, bold=True)
            tp.alignment = PP_ALIGN.CENTER

        # Event text
        if event:
            te = slide.shapes.add_textbox(x - Inches(0.8), line_y + Inches(0.3), Inches(1.6), Inches(1.5))
            te.text_frame.word_wrap = True
            ep = te.text_frame.paragraphs[0]
            _run(ep, event, size=Pt(11), color=theme.text_color, bold=True)
            if brief:
                ep2 = te.text_frame.add_paragraph()
                _run(ep2, brief, size=Pt(9), color=theme.secondary_color)


def _comparison(prs, info, theme):
    """Side-by-side comparison with VS divider."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = info.get("title", "Comparison")
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11.5), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    _run(p, title, size=Pt(26), color=theme.title_color, bold=True)
    _thin_line(slide, 0.8, 0.95, 3.0, theme.accent_bar_color)

    left_data = info.get("left", {})
    right_data = info.get("right", {})
    vs_label = info.get("vs_label", "VS")

    # Left panel
    _comp_panel(slide, Inches(0.5), Inches(1.3), Inches(5.2),
                left_data.get("label", "Option A"),
                left_data.get("points", []), theme,
                left_data.get("icon", "A"))

    # Right panel
    _comp_panel(slide, Inches(7.6), Inches(1.3), Inches(5.2),
                right_data.get("label", "Option B"),
                right_data.get("points", []), theme,
                right_data.get("icon", "B"))

    # VS badge in the center
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(6.15), Inches(3.3), Inches(1.0), Inches(1.0),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = _hex(theme.accent_bar_color)
    badge.line.fill.background()
    # VS text
    vstb = slide.shapes.add_textbox(Inches(6.15), Inches(3.5), Inches(1.0), Inches(0.6))
    vsp = vstb.text_frame.paragraphs[0]
    _run(vsp, vs_label, size=Pt(18), color="FFFFFF", bold=True)
    vsp.alignment = PP_ALIGN.CENTER


def _comp_panel(slide, x, y, w, label, points, theme, icon):
    """Helper: draw one comparison panel."""
    # Background card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(5.2))
    card.fill.solid()
    card.fill.fore_color.rgb = _hex("F5F7FA" if theme.background_color == "FFFFFF" else theme.background_color)
    card.line.color.rgb = _hex(theme.secondary_color)
    card.line.width = Pt(0.5)

    # Label
    lb = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), w - Inches(0.6), Inches(0.4))
    lp = lb.text_frame.paragraphs[0]
    _run(lp, f"{icon}  {label}", size=Pt(16), color=theme.title_color, bold=True)

    # Points
    pb = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.8), w - Inches(0.6), Inches(4.0))
    pb.text_frame.word_wrap = True
    for j, pt in enumerate(points):
        text = pt.get("text", "") if isinstance(pt, dict) else str(pt)
        pp = pb.text_frame.paragraphs[0] if j == 0 else pb.text_frame.add_paragraph()
        _run(pp, f"  {text}", size=Pt(13), color=theme.text_color)
        pp.space_after = Pt(8)
        pp.space_before = Pt(4)


def _big_idea(prs, info, theme):
    """One massive statement — like a billboard. Minimal design, maximum impact."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Full color background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5),
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = _hex(theme.primary_color)
    bg.line.fill.background()

    # Decorative top line
    _thin_line(slide, 5.0, 2.4, 3.333, theme.accent_bar_color)

    # Big text
    big_text = info.get("big_text", info.get("title", ""))
    tb = slide.shapes.add_textbox(Inches(1.5), Inches(2.7), Inches(10.333), Inches(2.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    _run(p, big_text, size=Pt(38), color="FFFFFF", bold=True)
    p.alignment = PP_ALIGN.CENTER

    # Sub text
    sub = info.get("sub_text", info.get("subtitle", ""))
    if sub:
        _thin_line(slide, 5.0, 5.1, 3.333, theme.accent_bar_color)
        tb2 = slide.shapes.add_textbox(Inches(2.0), Inches(5.3), Inches(9.333), Inches(0.8))
        tb2.text_frame.word_wrap = True
        p2 = tb2.text_frame.paragraphs[0]
        _run(p2, sub, size=Pt(18), color="B0C4DE")
        p2.alignment = PP_ALIGN.CENTER


# ============================================================================
# DECORATIONS — safe thin shapes only
# ============================================================================

def _top_bar(slide, theme):
    """Thin colored bar across the top of the slide."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(0.04),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(theme.accent_bar_color)
    shape.line.fill.background()


def _bottom_bar(slide, theme):
    """Subtle line at the bottom of content area."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(7.15), Inches(11.6), Inches(0.015),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(theme.secondary_color)
    shape.line.fill.background()


def _slide_num(slide, num, theme):
    """Slide number at bottom-right corner."""
    tb = slide.shapes.add_textbox(Inches(12.0), Inches(7.1), Inches(1.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    _run(p, str(num), size=Pt(8), color=theme.secondary_color)
    p.alignment = PP_ALIGN.RIGHT


def _left_stripe(slide, theme):
    """Left-side vertical accent stripe for section slides."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(0.06), Inches(7.5),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(theme.accent_bar_color)
    shape.line.fill.background()


def _slide_progress(slide, current, total, theme):
    """Tiny progress dots at the very bottom edge."""
    dot_size = Inches(0.04)
    gap = Inches(0.08)
    total_w = total * (dot_size + gap) - gap
    start_x = Inches((13.333 - Inches(total_w).inches) / 2)
    dot_y = Inches(7.35)

    for idx in range(total):
        x = start_x + idx * (dot_size + gap)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, dot_y, dot_size, dot_size)
        dot.fill.solid()
        if idx == current:
            dot.fill.fore_color.rgb = _hex(theme.accent_bar_color)
        else:
            dot.fill.fore_color.rgb = _hex("CCCCCC")
        dot.line.fill.background()


def _notes(slide, info):
    """Add speaker notes safely."""
    text = info.get("notes", "")
    if text:
        try:
            if slide.has_notes_slide:
                slide.notes_slide.notes_text_frame.text = text
        except Exception:
            pass
